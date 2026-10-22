"""Gera docs/slides/apresentacao_modificacao.pptx (identidade visual do deck da equipe).

Uso:  py -3.12 docs/slides/build_slides.py      (precisa de python-pptx)
"""
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.dml import MSO_LINE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
OUT = os.path.join(HERE, "apresentacao_modificacao.pptx")

# Paleta extraída dos slides da equipe
BG = RGBColor(0x01, 0x13, 0x2B)
CARD = RGBColor(0x0D, 0x26, 0x46)
LINE = RGBColor(0x1D, 0x2D, 0x42)
YELLOW = RGBColor(0xF1, 0xB0, 0x21)
WHITE = RGBColor(0xF7, 0xF7, 0xF7)
MUTED = RGBColor(0xA9, 0xB6, 0xC8)
BLUE = RGBColor(0x53, 0x86, 0xC4)
AG = [RGBColor(0x2B, 0x7B, 0xBA), RGBColor(0xE8, 0x86, 0x2A), RGBColor(0x3A, 0x9E, 0x5C)]
FONT = "Segoe UI"
FONT_BOLD = "Segoe UI Semibold"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
N = [0]


def text(slide, x, y, w, h, runs, size=20, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=None, spacing=1.1):
    """runs: str | list de parágrafos; cada parágrafo str ou lista de (texto, {opções})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for seg in (para if isinstance(para, list) else [(para, {})]):
            t, o = seg if isinstance(seg, tuple) else (seg, {})
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.name = o.get("font", font or (FONT_BOLD if o.get("bold", bold) else FONT))
            f.color.rgb = o.get("color", color)
    return tb


def rect(slide, x, y, w, h, fill=CARD, shape=MSO_SHAPE.ROUNDED_RECTANGLE, line=None, radius=0.06):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1.5)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def circle(slide, cx, cy, r, fill, line=None):
    s = rect(slide, cx - r, cy - r, 2 * r, 2 * r, fill=fill, shape=MSO_SHAPE.OVAL, line=line)
    return s


def connector(slide, x1, y1, x2, y2, color=MUTED, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def target(slide, cx, cy, size=0.22):
    halo = circle(slide, cx, cy, size * 1.25, LINE)
    x = slide.shapes.add_shape(MSO_SHAPE.MATH_MULTIPLY, Inches(cx - size), Inches(cy - size),
                               Inches(2 * size), Inches(2 * size))
    x.fill.solid(); x.fill.fore_color.rgb = MUTED; x.line.fill.background(); x.shadow.inherit = False
    return halo


def new_slide(title=None, notes=""):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    N[0] += 1
    if title:
        rect(s, 0.42, 0.42, 1.95, 0.07, fill=YELLOW, shape=MSO_SHAPE.RECTANGLE)
        text(s, 0.42, 0.62, 11.3, 1.0, title, size=36, bold=True)
    s.shapes.add_picture(os.path.join(HERE, "logo.png"), Inches(12.5), Inches(0.26), Inches(0.7), Inches(0.7))
    num = circle(s, 12.88, 7.03, 0.23, BLUE)
    tf = num.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(N[0]); r.font.size = Pt(12); r.font.color.rgb = WHITE; r.font.name = FONT
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def three_agents_one_net(s, cx, top, shared):
    """Diagrama: 3 agentes ligados a 1 rede (shared) ou cada um à sua."""
    xs = [cx - 1.35, cx, cx + 1.35]
    ay = top + 2.25
    if shared:
        net = rect(s, cx - 0.95, top, 1.9, 0.8, fill=YELLOW, radius=0.25)
        nets = [(cx, top + 0.8)] * 3
        text(s, cx - 0.95, top, 1.9, 0.8, "1 rede", size=20, bold=True, color=BG,
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    else:
        nets = []
        for x in xs:
            rect(s, x - 0.55, top, 1.1, 0.8, fill=CARD, line=MUTED, radius=0.25)
            text(s, x - 0.55, top, 1.1, 0.8, "rede", size=16, bold=True, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.MIDDLE)
            nets.append((x, top + 0.8))
    for (nx, ny), x, c in zip(nets, xs, AG):
        connector(s, nx, ny, x, ay - 0.32, color=MUTED, width=2.25)
        circle(s, x, ay, 0.32, c)


# ---------------------------------------------------------------- 1. capa
s = new_slide(notes=(
    "Boa tarde. Eu vou apresentar a parte prática do nosso miniprojeto: o que nós executamos do artigo, "
    "como os nossos números se comparam com os dele e a modificação que propusemos."))
rect(s, 0.9, 2.45, 1.95, 0.07, fill=YELLOW, shape=MSO_SHAPE.RECTANGLE)
text(s, 0.9, 2.7, 10.5, 1.9, ["Um cérebro para três agentes"], size=54, bold=True)
text(s, 0.9, 3.85, 8.2, 1.2,
     "Reprodução do artigo da MARLlib e o efeito de compartilhar a rede neural entre os agentes",
     size=22, color=MUTED)
text(s, 0.9, 5.2, 10.5, 0.5, "Artigo base: Hu et al., MARLlib, Journal of Machine Learning Research, v. 24, 2023",
     size=14, color=MUTED)
for i, (dx, dy) in enumerate([(10.6, 5.7), (11.5, 5.0), (12.1, 6.1)]):
    circle(s, dx, dy, 0.28, AG[i])
for dx, dy in [(10.1, 4.9), (11.1, 6.5), (12.4, 4.8)]:
    target(s, dx, dy, 0.16)

# ---------------------------------------------------------------- 2. o desafio
s = new_slide("O desafio", notes=(
    "O cenário que usamos é o simple_spread. São três agentes e três alvos num plano. O objetivo do time é "
    "cobrir os três alvos sem colidir. O detalhe importante: ninguém diz qual agente vai para qual alvo. "
    "Cada agente só vê o mundo a partir da sua posição, e eles precisam aprender a se dividir sozinhos. "
    "A pontuação é negativa: soma a distância de cada alvo até o agente mais próximo, mais as colisões. "
    "Quanto mais perto de zero, melhor."))
rect(s, 0.55, 1.75, 5.3, 4.95, fill=CARD, radius=0.04)
pts_t = [(1.7, 2.8), (4.55, 3.15), (2.9, 5.55)]
pts_a = [(1.35, 4.6), (3.6, 2.45), (4.85, 5.1)]
for (tx, ty), (ax, ay), c in zip([pts_t[0], pts_t[1], pts_t[2]], [pts_a[0], pts_a[1], pts_a[2]], AG):
    connector(s, ax, ay, tx, ty, color=c, width=1.5).line.dash_style = MSO_LINE.DASH
for tx, ty in pts_t:
    target(s, tx, ty)
for (ax, ay), c in zip(pts_a, AG):
    circle(s, ax, ay, 0.3, c)
items = [
    ("3 agentes, 3 alvos", "Cobrir todos os alvos sem colidir."),
    ("Ninguém tem alvo fixo", "Eles precisam aprender a se dividir."),
    ("Cada um vê só a sua volta", "Posição dos alvos e dos colegas em relação a ele."),
]
for i, (h, b) in enumerate(items):
    y = 1.85 + i * 1.3
    text(s, 6.45, y, 0.6, 0.8, str(i + 1), size=34, bold=True, color=YELLOW)
    text(s, 7.15, y + 0.02, 5.6, 1.2, [[(h, {"bold": True, "size": 24})], [(b, {"size": 18, "color": MUTED})]])
rect(s, 6.45, 5.75, 6.4, 0.95, fill=LINE, radius=0.12)
text(s, 6.7, 5.75, 6.0, 0.95,
     [[("Pontuação: ", {"bold": True, "color": YELLOW}),
       ("– distância de cada alvo ao agente mais próximo – colisões. Mais perto de zero = melhor.", {})]],
     size=16, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 3. o artigo
s = new_slide("O que o artigo propõe", notes=(
    "O artigo apresenta a MARLlib, uma biblioteca que funciona como uma tomada universal: conecta qualquer "
    "ambiente multiagente a qualquer algoritmo. Ela tem três mecanismos. O terceiro decide quais agentes "
    "compartilham a mesma rede neural. O artigo oferece essas opções, mas em nenhum experimento mede qual "
    "delas é melhor. É exatamente aí que entra a nossa modificação."))
text(s, 0.42, 1.55, 12, 0.6,
     "MARLlib: uma “tomada universal” que conecta qualquer ambiente multiagente a qualquer algoritmo.",
     size=20, color=MUTED)
cards = [("1", "Formato único", "Todo ambiente entrega os dados do mesmo jeito."),
         ("2", "Dados por agente", "Algoritmos de famílias diferentes no mesmo código."),
         ("3", "Quem compartilha a rede", "Todos, por grupo ou cada um com a sua.")]
for i, (n, h, b) in enumerate(cards):
    x = 0.42 + i * 4.2
    hi = i == 2
    rect(s, x, 2.45, 3.95, 2.75, fill=YELLOW if hi else CARD, radius=0.05)
    fg = BG if hi else WHITE
    text(s, x + 0.35, 2.65, 1, 0.9, n, size=40, bold=True, color=BG if hi else YELLOW)
    text(s, x + 0.35, 3.55, 3.3, 1.6,
         [[(h, {"bold": True, "size": 22, "color": fg})], [(b, {"size": 16, "color": BG if hi else MUTED})]])
text(s, 8.82, 5.4, 4.1, 0.9, [[("O artigo oferece, mas não mede. ", {"bold": True, "color": YELLOW}),
                               ("É aqui que entramos.", {})]], size=18)

# ---------------------------------------------------------------- 4. reprodução
s = new_slide("Executamos o que está no artigo", notes=(
    "Primeiro reproduzimos o artigo: código oficial, hiperparâmetros oficiais, dois algoritmos, IPPO e MAPPO, "
    "três sementes cada. O que confirmamos: os dois empatam, como no artigo. O que não bateu: ficamos cerca de "
    "30 pontos abaixo. Treinamos quatro vezes mais, 2 milhões de passos, e a diferença caiu para 22. "
    "Então parte é falta de treino, e o resto vem de detalhes que o artigo não publica. É um achado de "
    "reprodutibilidade, não um erro de execução."))
text(s, 0.42, 1.55, 12, 0.6, "Código oficial · hiperparâmetros oficiais · IPPO e MAPPO · 3 sementes · 500 mil passos",
     size=18, color=MUTED)
for i, (alg, nos, art) in enumerate([("IPPO", "−96,4", "−65,7"), ("MAPPO", "−97,2", "−66,3")]):
    x = 0.42 + i * 3.35
    rect(s, x, 2.35, 3.1, 3.0, fill=CARD, radius=0.05)
    text(s, x + 0.3, 2.55, 2.5, 0.5, alg, size=22, bold=True, color=YELLOW)
    text(s, x + 0.3, 3.15, 2.6, 1.0, [[(nos, {"size": 40, "bold": True})], [("nós", {"size": 14, "color": MUTED})]])
    text(s, x + 0.3, 4.3, 2.6, 1.0, [[(art, {"size": 26, "bold": True, "color": MUTED})],
                                      [("artigo", {"size": 14, "color": MUTED})]])
rows = [("✔", "Confirmado", "IPPO e MAPPO empatam, como no artigo.", WHITE),
        ("✖", "Não bateu", "Ficamos ~30 pontos abaixo do artigo.", WHITE),
        ("→", "Treinamos 4× mais", "Com 2 milhões de passos, a diferença cai para ~22.", YELLOW)]
for i, (ic, h, b, c) in enumerate(rows):
    y = 2.4 + i * 1.0
    text(s, 7.35, y, 0.5, 0.8, ic, size=26, bold=True, color=YELLOW)
    text(s, 7.95, y + 0.02, 5.0, 0.9, [[(h, {"bold": True, "size": 20, "color": c})],
                                        [(b, {"size": 16, "color": MUTED})]])
rect(s, 7.35, 5.55, 5.5, 0.95, fill=YELLOW, radius=0.12)
text(s, 7.6, 5.55, 5.1, 0.95, "O artigo não publica a configuração completa: achado de reprodutibilidade.",
     size=16, bold=True, color=BG, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 5. vídeo
s = new_slide("Os agentes antes e depois do treino", notes=(
    "Aqui estão os agentes reais do nosso treino, na mesma situação inicial. À esquerda, sem treino: andam ao "
    "acaso e se trombam. No meio, depois de 160 mil passos. À direita, depois de 2 milhões: cada agente vai "
    "quase em linha reta para um alvo diferente. Ninguém programou essa divisão, eles aprenderam. No fim, a "
    "pontuação média em 200 episódios."))
vw = 12.0
vh = vw * 594 / 1408
mv = s.shapes.add_movie(os.path.join(ROOT, "results", "video", "agentes_treinando.mp4"),
                        Inches((13.333 - vw) / 2), Inches(1.75), Inches(vw), Inches(vh),
                        poster_frame_image=os.path.join(HERE, "poster.png"), mime_type="video/mp4")

# ---------------------------------------------------------------- 6. a modificação
s = new_slide("Nossa modificação: um cérebro ou três?", notes=(
    "A nossa pergunta foi: agentes iguais precisam de cérebros separados? Testamos as duas opções. "
    "À esquerda, o padrão do artigo: os três agentes usam a mesma rede neural. À direita, cada agente tem a "
    "sua. Na prática foi um único parâmetro da biblioteca. Todo o resto ficou idêntico, então qualquer "
    "diferença vem só disso. Medimos duas coisas: pontuação e tempo de treino."))
for i, (shared, h, sub, code) in enumerate([
        (True, "Rede compartilhada", "padrão do artigo", 'share_policy = "all"'),
        (False, "Uma rede por agente", "variação", 'share_policy = "individual"')]):
    x = 0.42 + i * 6.35
    rect(s, x, 1.65, 6.1, 4.45, fill=CARD, radius=0.04)
    text(s, x + 0.35, 1.85, 5.4, 0.9, [[(h, {"bold": True, "size": 24})], [(sub, {"size": 15, "color": MUTED})]])
    three_agents_one_net(s, x + 3.05, 2.85, shared)
    rect(s, x + 0.35, 5.45, 5.4, 0.5, fill=BG, radius=0.2)
    text(s, x + 0.35, 5.45, 5.4, 0.5, code, size=15, font="Consolas", color=YELLOW if shared else MUTED,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.42, 6.35, 12.4, 0.6, [[("Um único parâmetro. ", {"bold": True, "color": YELLOW}),
                                 ("Todo o resto igual: mesmo ambiente, mesmos hiperparâmetros, 3 sementes, 500 mil passos.",
                                  {"color": MUTED})]], size=17)

# ---------------------------------------------------------------- 7. resultado
s = new_slide("O resultado", notes=(
    "O resultado. Com a rede compartilhada, o treino levou 17 minutos em vez de 30: 42 por cento menos tempo. "
    "E a pontuação não piorou: ficou igual ou um pouco melhor, com menos variação entre as sementes. "
    "Para ver se isso dependia do meu computador, um colega rodou o mesmo teste numa máquina diferente, "
    "um Ryzen, e o ganho foi de 33 por cento. O valor muda com o hardware, mas a tendência se mantém."))
text(s, 0.42, 1.75, 4.2, 1.6, "−42%", size=88, bold=True, color=YELLOW)
text(s, 0.42, 3.25, 4.0, 0.9, "de tempo de treino com a rede compartilhada", size=20)
stats = [("Pontuação", "igual ou melhor", "IPPO −96,4 × −97,4   ·   MAPPO −97,2 × −98,7"),
         ("Outro computador", "−33%", "Ryzen 3 3300X: o ganho se repete")]
for i, (h, v, b) in enumerate(stats):
    y = 4.35 + i * 1.15
    rect(s, 0.42, y, 4.4, 1.0, fill=CARD, radius=0.1)
    text(s, 0.65, y + 0.1, 4.0, 0.9, [[(h + ": ", {"size": 15, "color": MUTED}), (v, {"bold": True, "size": 18})],
                                       [(b, {"size": 12, "color": MUTED})]])
cd = CategoryChartData()
cd.categories = ["IPPO", "MAPPO"]
cd.add_series("Rede compartilhada", (17.0, 17.8))
cd.add_series("Uma rede por agente", (28.7, 30.5))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5.3), Inches(1.65), Inches(7.6), Inches(5.0), cd)
ch = gf.chart
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(15); ch.legend.font.color.rgb = WHITE; ch.legend.font.name = FONT
ch.font.name = FONT
for ser, col in zip(ch.plots[0].series, (YELLOW, BLUE)):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = col
ch.plots[0].gap_width = 70
ch.plots[0].overlap = -8
pl = ch.plots[0]
pl.has_data_labels = True
dl = pl.data_labels
dl.number_format = '0.0" min"'; dl.number_format_is_linked = False
dl.position = XL_LABEL_POSITION.OUTSIDE_END
dl.font.size = Pt(16); dl.font.bold = True; dl.font.color.rgb = WHITE
ca, va = ch.category_axis, ch.value_axis
ca.tick_labels.font.size = Pt(18); ca.tick_labels.font.color.rgb = WHITE
ca.format.line.color.rgb = LINE
va.visible = False
va.has_major_gridlines = False
va.maximum_scale = 36; va.minimum_scale = 0
text(s, 5.3, 6.6, 7.6, 0.4, "Tempo médio por execução (min) · 500 mil passos · 3 sementes", size=13,
     color=MUTED, align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------- 8. por que funciona
s = new_slide("Por que funciona", notes=(
    "Por que isso funciona? Três motivos. Primeiro: com uma rede só, o computador atualiza uma rede por rodada, "
    "não três. Segundo: a rede compartilhada aprende com a experiência dos três agentes, três vezes mais "
    "exemplos, por isso o treino fica mais estável. Terceiro: os agentes têm o mesmo papel. Eles não fazem a "
    "mesma coisa porque cada um vê o mundo de um lugar diferente. É como três entregadores seguindo o mesmo "
    "manual: cada um aplica onde está."))
reasons = [("1", "Menos trabalho", "1 rede para atualizar por rodada, não 3."),
           ("2", "Mais experiência", "A rede aprende com as jogadas dos 3 agentes: 3× mais exemplos."),
           ("3", "Mesmo papel", "Os agentes agem diferente porque cada um vê o mundo de um lugar.")]
for i, (n, h, b) in enumerate(reasons):
    x = 0.42 + i * 4.2
    rect(s, x, 1.8, 3.95, 3.3, fill=CARD, radius=0.05)
    text(s, x + 0.35, 2.0, 1, 0.9, n, size=44, bold=True, color=YELLOW)
    text(s, x + 0.35, 3.0, 3.3, 2.0, [[(h, {"bold": True, "size": 24})], [(b, {"size": 17, "color": MUTED})]])
rect(s, 0.42, 5.45, 12.4, 1.05, fill=LINE, radius=0.12)
text(s, 0.75, 5.45, 11.8, 1.05, [[("Analogia: ", {"bold": True, "color": YELLOW}),
                                  ("três entregadores seguindo o mesmo manual. Escrever um manual para cada um dá 3× "
                                   "mais trabalho, e cada manual aprende com menos experiência.", {})]],
     size=18, anchor=MSO_ANCHOR.MIDDLE)

# ---------------------------------------------------------------- 9. conclusão
s = new_slide("O que isso muda no mundo real", notes=(
    "O que isso muda fora da sala de aula: numa frota de cem drones ou robôs iguais, dá para usar uma rede só, "
    "em vez de cem. O custo de treino deixa de crescer com o número de agentes. E treinar mais barato significa "
    "fazer mais experimentos no mesmo prazo. O limite: se os agentes tiverem papéis diferentes, como um drone "
    "de câmera e um de carga, o indicado é compartilhar por grupo, uma rede por tipo. Todo o código e os "
    "resultados estão no nosso repositório. Obrigado."))
pts = [("Frotas de agentes iguais", "100 drones ou robôs → 1 rede, não 100."),
       ("Treino mais barato", "Menos tempo de máquina = mais experimentos no mesmo prazo."),
       ("Limite", "Papéis diferentes (câmera × carga) → compartilhar por grupo.")]
for i, (h, b) in enumerate(pts):
    y = 1.8 + i * 1.3
    rect(s, 0.42, y + 0.12, 0.09, 0.85, fill=YELLOW if i < 2 else MUTED, shape=MSO_SHAPE.RECTANGLE)
    text(s, 0.8, y, 7.5, 1.2, [[(h, {"bold": True, "size": 24})], [(b, {"size": 18, "color": MUTED})]])
rect(s, 8.6, 1.8, 4.25, 3.75, fill=YELLOW, radius=0.05)
text(s, 8.95, 2.05, 3.6, 3.3, [[("Em uma frase", {"size": 16, "bold": True, "color": BG})],
                               [("Agentes iguais podem dividir um só cérebro: mesmo desempenho, ~40% menos tempo.",
                                 {"size": 24, "bold": True, "color": BG})]], spacing=1.15)
text(s, 0.42, 6.15, 12.4, 0.5, [[("Código e resultados: ", {"color": MUTED}),
                                 ("github.com/heitorfariass/miniprojeto-ia-marllib", {"bold": True})]], size=16)

prs.save(OUT)
print("ok", OUT, len(prs.slides), "slides")
